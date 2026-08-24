"""Dump the deck's slide titles and speaker notes to markdown.

The notes carry the numbers and the argument for each slide, so this file doubles
as the study document for the question-preparation sessions.

Run from repo root with the venv active:
    PYTHONIOENCODING=utf-8 python code/defense/dump_notes.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

DECK = Path("thesis/defense/ISFS_defense_V2.pptx")
OUT = Path("thesis/defense/speaker_notes.md")


def slide_text(slide):
    """Best-effort title: the first non-empty text line on the slide."""
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            return shape.text_frame.text.strip().splitlines()[0]
    return "(no text)"


def main():
    prs = Presentation(DECK)
    lines = [f"# Speaker notes: {DECK.name}",
             "",
             f"Generated from the deck by `code/defense/dump_notes.py`. "
             f"{len(prs.slides)} slides.",
             "",
             "Edit the notes in `code/defense/build_deck.py` and rebuild, not here.",
             ""]
    for i, slide in enumerate(prs.slides, start=1):
        title = slide_text(slide)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        lines.append(f"## {i}. {title}")
        lines.append("")
        lines.append(notes if notes else "_(no notes)_")
        lines.append("")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Saved: {OUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
